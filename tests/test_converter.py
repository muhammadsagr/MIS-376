"""اختبارات البرنامج:  python -m unittest discover -s tests"""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Cm, Pt

from orgchart.converter import convert
from orgchart.extractor import ExtractOptions, _extract_smartart, extract_from_presentation
from orgchart.hierarchy import build_hierarchy
from orgchart.models import ExtractionResult
from orgchart.textparse import parse_box_text
from samples.make_grades_sample import build as build_grades
from samples.make_sample import build as build_sample

DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _add_box(slide, text, left, top, width=5.0, height=2.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(left), Cm(top),
                                   Cm(width), Cm(height))
    frame = shape.text_frame
    lines = text.split("\n")
    frame.text = lines[0]
    for line in lines[1:]:
        para = frame.add_paragraph()
        para.text = line
        para.font.size = Pt(10)
    return shape


class TextParsingTests(unittest.TestCase):
    def test_multi_line_box(self):
        fields = parse_box_text("Ali Ahmed\nSales Manager\nSales Department\next 100")
        self.assertEqual(fields["name"], "Ali Ahmed")
        self.assertEqual(fields["title"], "Sales Manager")
        self.assertEqual(fields["department"], "Sales Department")
        self.assertEqual(fields["extra"], "ext 100")

    def test_inline_separator(self):
        fields = parse_box_text("منى صالح - مدير الموارد البشرية")
        self.assertEqual(fields["name"], "منى صالح")
        self.assertEqual(fields["title"], "مدير الموارد البشرية")

    def test_single_line_title_only(self):
        self.assertEqual(parse_box_text("المدير العام")["title"], "المدير العام")
        self.assertEqual(parse_box_text("Sara Ali")["name"], "Sara Ali")


class ConnectorChartTests(unittest.TestCase):
    """الحالة الشائعة: مربعات + خطوط ربط."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.pptx = build_sample(os.path.join(cls.tmp, "sample.pptx"))

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_hierarchy_and_excel(self):
        xlsx, nodes, result = convert(self.pptx, os.path.join(self.tmp, "out.xlsx"))
        self.assertTrue(os.path.isfile(xlsx))
        self.assertEqual(len(nodes), 7)

        roots = [n for n in nodes if n.parent_id is None]
        self.assertEqual(len(roots), 1)
        self.assertEqual(roots[0].name, "أحمد السالم")
        self.assertEqual(roots[0].direct_reports, 3)
        self.assertEqual(roots[0].total_reports, 6)
        self.assertEqual(max(n.level for n in nodes), 3)

        by_name = {n.name: n for n in nodes}
        self.assertEqual(by_name["فهد الشمري"].parent_id, by_name["منى العتيبي"].node_id)
        self.assertEqual(by_name["ريم الدوسري"].level, 3)

        wb = load_workbook(xlsx)
        self.assertEqual(len(wb.sheetnames), 3)
        sheet = wb[wb.sheetnames[0]]
        self.assertEqual(sheet.max_row, 8)                      # 7 rows + header
        headers = [c.value for c in sheet[1]]
        self.assertIn("المسمى الوظيفي", headers)          # الوضع الافتراضي: هيكل وظائف
        manager_col = headers.index("الوظيفة الأعلى") + 1
        managers = [sheet.cell(row=r, column=manager_col).value for r in range(2, 9)]
        self.assertIn("الرئيس التنفيذي", managers)

    def test_positions_vs_employees(self):
        """وضع الوظائف (افتراضي) مقابل وضع الموظفين."""
        _, jobs, _ = convert(self.pptx, os.path.join(self.tmp, "jobs.xlsx"))
        top = [n for n in jobs if n.parent_id is None][0]
        self.assertEqual(top.title, "الرئيس التنفيذي")     # الوظيفة هي التسمية
        self.assertEqual(top.name, "أحمد السالم")          # الشاغل يبقى في عمود منفصل
        self.assertEqual(top.display_name, "الرئيس التنفيذي")
        self.assertEqual(jobs[1].path, "الرئيس التنفيذي > المدير المالي")

        xlsx, staff, _ = convert(self.pptx, os.path.join(self.tmp, "staff.xlsx"),
                                 positions=False)
        top2 = [n for n in staff if n.parent_id is None][0]
        self.assertEqual(top2.display_name, "أحمد السالم")
        self.assertIn("الاسم", [c.value for c in load_workbook(xlsx)[
            load_workbook(xlsx).sheetnames[0]][1]])

    def test_slide_filter(self):
        _, nodes, _ = convert(self.pptx, os.path.join(self.tmp, "s1.xlsx"), slides=[1])
        self.assertEqual(len(nodes), 7)
        with self.assertRaises(ValueError):        # الشريحة رقم 2 غير موجودة
            convert(self.pptx, os.path.join(self.tmp, "s2.xlsx"), slides=[2])

    def test_slide_title_is_skipped_by_default(self):
        _, nodes, _ = convert(self.pptx, os.path.join(self.tmp, "t1.xlsx"))
        self.assertNotIn("الهيكل الوظيفي - شركة نموذجية", [n.raw_text for n in nodes])
        _, with_title, _ = convert(self.pptx, os.path.join(self.tmp, "t2.xlsx"),
                                   include_titles=True)
        self.assertEqual(len(with_title), 8)
        self.assertEqual(nodes[0].slide_title, "الهيكل الوظيفي - شركة نموذجية")


class GeometryFallbackTests(unittest.TestCase):
    """مربعات بدون خطوط ربط: يُستنتج التسلسل من المواقع."""

    def test_levels_from_positions(self):
        tmp = tempfile.mkdtemp()
        try:
            prs = Presentation()
            prs.slide_width, prs.slide_height = Cm(25.4), Cm(19.05)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            _add_box(slide, "Nora Ali\nGeneral Manager", 10.0, 1.0)
            _add_box(slide, "Omar Zaid\nFinance Manager", 4.0, 6.0)
            _add_box(slide, "Lina Adel\nIT Manager", 16.0, 6.0)
            _add_box(slide, "Sami Nabil\nAccountant", 4.0, 11.0)
            path = os.path.join(tmp, "geom.pptx")
            prs.save(path)

            _, nodes, result = convert(path, os.path.join(tmp, "geom.xlsx"))
            by_name = {n.name: n for n in nodes}
            self.assertEqual(len(nodes), 4)
            self.assertIsNone(by_name["Nora Ali"].parent_id)
            self.assertEqual(by_name["Omar Zaid"].parent_id, by_name["Nora Ali"].node_id)
            self.assertEqual(by_name["Sami Nabil"].parent_id, by_name["Omar Zaid"].node_id)
            self.assertTrue(any("مواقع المربعات" in w for w in result.warnings))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class GradeLayoutTests(unittest.TestCase):
    """هيكل وظائف بدرجات: خلية درجة + خلية مسمى، وعمود صفوف تحت كل مدير."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.pptx = os.path.join(cls.tmp, "grades.pptx")
        build_grades(cls.pptx)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_grade_cells_and_flat_columns(self):
        _, nodes, _ = convert(self.pptx, os.path.join(self.tmp, "grades.xlsx"))
        by_title = {n.title: n for n in nodes if n.title}

        gm = by_title["General Manager of Human Capital Services"]
        self.assertEqual(gm.grade, "M5")                 # دُمجت خلية الدرجة مع المسمى
        self.assertIsNone(gm.parent_id)
        self.assertEqual(gm.direct_reports, 3)           # ثلاثة مدراء

        director = by_title["Payroll and Benefits Director"]
        self.assertEqual(director.grade, "M4")
        self.assertEqual(director.parent_id, gm.node_id)
        self.assertEqual(director.direct_reports, 5)     # الصفوف تتبع المدير مباشرة
        self.assertEqual(max(n.level for n in nodes), 3)  # ولا تتسلسل تحت بعضها

        bare = [n for n in nodes if not n.title and n.grade]
        self.assertEqual(len(bare), 14)                  # صفوف الدرجات بلا مسمى
        self.assertTrue(all(n.level == 3 for n in bare))
        self.assertEqual(bare[0].display_name, f"وظيفة بدرجة {bare[0].grade}")

    def test_grade_table_layout(self):
        """نفس الهيكل لكنه مرسوم كجداول (عمود درجة + عمود مسمى)."""
        prs = Presentation()
        prs.slide_width, prs.slide_height = Cm(33.87), Cm(19.05)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        blocks = [
            (1.5, [("M4", "Payroll and Benefits Director"), ("39", ""), ("38", ""), ("37", "")]),
            (14.0, [("M3", "Employees Services Director"), ("38", ""), ("37", "")]),
        ]
        for left, rows in blocks:
            table = slide.shapes.add_table(len(rows), 2, Cm(left), Cm(3.0),
                                           Cm(11.0), Cm(1.0 * len(rows))).table
            for r, (grade, title) in enumerate(rows):
                table.cell(r, 0).text = grade
                table.cell(r, 1).text = title
        path = os.path.join(self.tmp, "grade_table.pptx")
        prs.save(path)

        _, nodes, _ = convert(path, os.path.join(self.tmp, "grade_table.xlsx"))
        head = [n for n in nodes if n.title == "Payroll and Benefits Director"][0]
        self.assertEqual(head.grade, "M4")
        self.assertEqual(head.direct_reports, 3)
        self.assertEqual([n.grade for n in nodes if n.parent_id == head.node_id],
                         ["39", "38", "37"])


class TableTests(unittest.TestCase):
    """جدول فيه عمود (المدير المباشر)."""

    def test_manager_column(self):
        tmp = tempfile.mkdtemp()
        try:
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            rows = [
                ["الاسم", "المسمى الوظيفي", "القسم", "المدير المباشر"],
                ["أحمد", "الرئيس التنفيذي", "الإدارة العليا", ""],
                ["منى", "مدير مالي", "المالية", "أحمد"],
                ["فهد", "محاسب", "المالية", "منى"],
            ]
            table = slide.shapes.add_table(len(rows), 4, Cm(1), Cm(1), Cm(20), Cm(6)).table
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
            path = os.path.join(tmp, "table.pptx")
            prs.save(path)

            _, nodes, _ = convert(path, os.path.join(tmp, "table.xlsx"))
            by_name = {n.name: n for n in nodes}
            self.assertEqual(len(nodes), 3)
            self.assertEqual(by_name["فهد"].level, 3)
            self.assertEqual(by_name["منى"].department, "المالية")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


class SmartArtTests(unittest.TestCase):
    """قراءة التسلسل من بيانات مخطط SmartArt مباشرة."""

    DATA = f"""<dgm:dataModel xmlns:dgm="{DGM}" xmlns:a="{A}">
      <dgm:ptLst>
        <dgm:pt modelId="1" type="doc"/>
        <dgm:pt modelId="10"><dgm:t><a:p><a:r><a:t>Aisha Omar</a:t></a:r></a:p>
          <a:p><a:r><a:t>CEO</a:t></a:r></a:p></dgm:t></dgm:pt>
        <dgm:pt modelId="20"><dgm:t><a:p><a:r><a:t>Bilal Said</a:t></a:r></a:p>
          <a:p><a:r><a:t>CFO</a:t></a:r></a:p></dgm:t></dgm:pt>
        <dgm:pt modelId="30"><dgm:t><a:p><a:r><a:t>Dana Fahd</a:t></a:r></a:p>
          <a:p><a:r><a:t>Accountant</a:t></a:r></a:p></dgm:t></dgm:pt>
        <dgm:pt modelId="99" type="pres"><dgm:t><a:p><a:r><a:t>ignore me</a:t></a:r></a:p></dgm:t></dgm:pt>
      </dgm:ptLst>
      <dgm:cxnLst>
        <dgm:cxn modelId="c1" srcId="1" destId="10" type="parOf" destOrd="0"/>
        <dgm:cxn modelId="c2" srcId="10" destId="20" type="parOf" destOrd="0"/>
        <dgm:cxn modelId="c3" srcId="20" destId="30" type="parOf" destOrd="0"/>
        <dgm:cxn modelId="c4" srcId="10" destId="99" type="presOf"/>
      </dgm:cxnLst>
    </dgm:dataModel>"""

    class _Part:
        def __init__(self, blob):
            self.blob = blob

    class _Shape:
        def __init__(self, element, part):
            self._element, self.part = element, part

    class _SlidePart:
        def __init__(self, blob):
            self._blob = blob

        def related_part(self, rid):
            assert rid == "rId2"
            return SmartArtTests._Part(self._blob)

    def test_diagram_hierarchy(self):
        frame = etree.fromstring(
            f'<p:graphicFrame xmlns:p="{P}" xmlns:dgm="{DGM}" xmlns:r="{R}">'
            f'<dgm:relIds r:dm="rId2"/></p:graphicFrame>')
        shape = self._Shape(frame, self._SlidePart(self.DATA.encode("utf-8")))

        result = ExtractionResult()
        added = _extract_smartart(shape, 1, "s1_", result, True)
        self.assertEqual(added, 3)                              # the "pres" point is skipped

        nodes = build_hierarchy(result)
        by_name = {n.name: n for n in nodes}
        self.assertEqual(by_name["Aisha Omar"].level, 1)
        self.assertEqual(by_name["Bilal Said"].parent_id, by_name["Aisha Omar"].node_id)
        self.assertEqual(by_name["Dana Fahd"].level, 3)
        self.assertEqual(by_name["Aisha Omar"].total_reports, 2)


class ErrorTests(unittest.TestCase):
    def test_missing_file(self):
        with self.assertRaises(FileNotFoundError):
            convert("no_such_file.pptx")

    def test_wrong_extension(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, "x.ppt")
            with open(path, "wb") as handle:
                handle.write(b"0")
            with self.assertRaises(ValueError):
                convert(path)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_empty_presentation(self):
        tmp = tempfile.mkdtemp()
        try:
            path = os.path.join(tmp, "empty.pptx")
            Presentation().save(path)
            with self.assertRaises(ValueError):
                convert(path)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
