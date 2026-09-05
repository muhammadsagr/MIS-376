"""Writes the extracted hierarchy into a formatted Excel workbook."""

from __future__ import annotations

from datetime import datetime
from typing import Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import ExtractionResult, OrgNode

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
LEVEL_FILLS = ["DDEBF7", "E2EFDA", "FFF2CC", "FCE4D6", "EDEDED", "F2F2F2"]
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS: Sequence[tuple] = (
    ("م", 6),
    ("الشريحة", 9),
    ("عنوان الشريحة", 22),
    ("المستوى", 9),
    ("الاسم", 26),
    ("المسمى الوظيفي", 26),
    ("القسم / الإدارة", 22),
    ("المدير المباشر", 26),
    ("مرؤوسون مباشرون", 16),
    ("إجمالي المرؤوسين", 16),
    ("المسار الوظيفي", 46),
    ("المصدر", 12),
    ("المعرّف", 12),
    ("النص الأصلي", 40),
)

SOURCE_LABEL = {"smartart": "SmartArt", "shape": "أشكال", "table": "جدول"}


def _style_header(ws: Worksheet, widths: Sequence[tuple], rtl: bool) -> None:
    for idx, (title, width) in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=idx, value=title)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.sheet_view.rightToLeft = rtl


def _write_main_sheet(ws: Worksheet, nodes: List[OrgNode], by_id: Dict[str, OrgNode],
                      rtl: bool) -> None:
    _style_header(ws, COLUMNS, rtl)
    for row_idx, node in enumerate(nodes, start=2):
        parent = by_id.get(node.parent_id) if node.parent_id else None
        values = [
            row_idx - 1,
            node.slide_index,
            node.slide_title,
            node.level,
            node.name,
            node.title,
            node.department,
            parent.display_name if parent else "",
            node.direct_reports,
            node.total_reports,
            node.path,
            SOURCE_LABEL.get(node.source, node.source),
            node.node_id,
            node.raw_text.replace("\n", " / "),
        ]
        fill = PatternFill("solid", fgColor=LEVEL_FILLS[min(node.level - 1, len(LEVEL_FILLS) - 1)])
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = BORDER
            cell.fill = fill
            cell.alignment = Alignment(
                vertical="center", wrap_text=col_idx in (3, 5, 6, 7, 8, 11, 14),
                horizontal="center" if col_idx in (1, 2, 4, 9, 10, 12) else "right" if rtl else "left",
            )
            if col_idx in (5, 6) and node.level == 1:
                cell.font = Font(bold=True)
    if len(nodes):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(nodes) + 1}"


def _write_tree_sheet(ws: Worksheet, nodes: List[OrgNode], rtl: bool) -> None:
    cols = (("المستوى", 9), ("الهيكل الشجري", 60), ("المسمى الوظيفي", 28),
            ("القسم / الإدارة", 22), ("عدد المرؤوسين", 14))
    _style_header(ws, cols, rtl)
    for row_idx, node in enumerate(nodes, start=2):
        indent = max(node.level - 1, 0)
        label = ("└─ " if indent else "") + (node.display_name or "-")
        ws.cell(row=row_idx, column=1, value=node.level).alignment = Alignment(horizontal="center")
        cell = ws.cell(row=row_idx, column=2, value=label)
        cell.alignment = Alignment(indent=indent * 2, horizontal="right" if rtl else "left")
        if node.level == 1:
            cell.font = Font(bold=True)
        ws.cell(row=row_idx, column=3, value=node.title)
        ws.cell(row=row_idx, column=4, value=node.department)
        ws.cell(row=row_idx, column=5, value=node.direct_reports).alignment = \
            Alignment(horizontal="center")
        for col in range(1, 6):
            ws.cell(row=row_idx, column=col).border = BORDER
        if node.level > 1:
            ws.row_dimensions[row_idx].outlineLevel = min(node.level - 1, 7)
    ws.sheet_properties.outlinePr.summaryBelow = False


def _write_summary_sheet(ws: Worksheet, nodes: List[OrgNode], result: ExtractionResult,
                         rtl: bool) -> None:
    _style_header(ws, (("البيان", 34), ("القيمة", 58)), rtl)
    levels = {}
    departments = {}
    for node in nodes:
        levels[node.level] = levels.get(node.level, 0) + 1
        if node.department:
            departments[node.department] = departments.get(node.department, 0) + 1

    rows = [
        ("ملف العرض التقديمي", result.source_file),
        ("تاريخ التحويل", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("عدد الشرائح المستخدمة", ", ".join(str(s) for s in result.slides_used()) or "-"),
        ("إجمالي عدد الوظائف", len(nodes)),
        ("عدد المستويات الإدارية", max(levels) if levels else 0),
        ("عدد الوظائف في القمة (بدون مدير)", sum(1 for n in nodes if n.parent_id is None)),
        ("عدد الوظائف بدون مرؤوسين", sum(1 for n in nodes if n.direct_reports == 0)),
    ]
    for level in sorted(levels):
        rows.append((f"عدد الوظائف في المستوى {level}", levels[level]))
    for dept, count in sorted(departments.items(), key=lambda kv: -kv[1]):
        rows.append((f"القسم: {dept}", count))
    for warning in result.warnings:
        rows.append(("ملاحظة", warning))

    for row_idx, (label, value) in enumerate(rows, start=2):
        c1 = ws.cell(row=row_idx, column=1, value=label)
        c2 = ws.cell(row=row_idx, column=2, value=value)
        c1.font = Font(bold=True)
        for cell in (c1, c2):
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True,
                                       horizontal="right" if rtl else "left")


def export_to_excel(nodes: List[OrgNode], result: ExtractionResult, output_path: str,
                    rtl: bool = True) -> str:
    """Create the workbook (3 sheets) and return the path it was saved to."""
    wb = Workbook()
    by_id = {n.node_id: n for n in nodes}

    main = wb.active
    main.title = "الهيكل الوظيفي"
    _write_main_sheet(main, nodes, by_id, rtl)

    _write_tree_sheet(wb.create_sheet("العرض الشجري"), nodes, rtl)
    _write_summary_sheet(wb.create_sheet("ملخص وتقرير"), nodes, result, rtl)

    wb.save(output_path)
    return output_path
