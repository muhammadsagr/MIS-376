"""ينشئ ملف PowerPoint تجريبي فيه هيكل وظيفي، لاختبار البرنامج بسرعة.

    python samples/make_sample.py            -> samples/sample_orgchart.pptx
"""

from __future__ import annotations

import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Cm, Pt

BOXES = [
    # (key, parent, name, title, left cm, top cm)
    ("ceo", None, "أحمد السالم", "الرئيس التنفيذي", 11.0, 1.5),
    ("cfo", "ceo", "منى العتيبي", "المدير المالي", 4.0, 6.0),
    ("cto", "ceo", "خالد الحربي", "مدير تقنية المعلومات", 11.0, 6.0),
    ("hrd", "ceo", "سارة القحطاني", "مدير الموارد البشرية", 18.0, 6.0),
    ("acc", "cfo", "فهد الشمري", "محاسب أول", 4.0, 10.5),
    ("dev", "cto", "ريم الدوسري", "مهندس برمجيات", 11.0, 10.5),
    ("rec", "hrd", "نورة الغامدي", "أخصائي توظيف", 18.0, 10.5),
]


def build(path: str) -> str:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(25.4), Cm(19.05)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "الهيكل الوظيفي - شركة نموذجية"

    shapes = {}
    for key, _parent, name, title, left, top in BOXES:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     Cm(left), Cm(top), Cm(5.5), Cm(2.2))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = name
        frame.paragraphs[0].runs[0].font.size = Pt(12)
        para = frame.add_paragraph()
        para.text = title
        para.font.size = Pt(10)
        shapes[key] = box

    for key, parent, *_ in BOXES:
        if not parent:
            continue
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.ELBOW, Cm(0), Cm(0), Cm(1), Cm(1))
        connector.begin_connect(shapes[parent], 2)   # bottom of the manager box
        connector.end_connect(shapes[key], 0)        # top of the employee box

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    prs.save(path)
    return path


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else \
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_orgchart.pptx")
    print("تم إنشاء الملف التجريبي:", build(target))
